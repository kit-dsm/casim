from __future__ import annotations

import math
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


SHEET_NAME = "Instances - Henn"
INSTANCE_RE = re.compile(r"H_(abc1|abc2|ran1|ran2)_(40|60|80|100)_\d+")
PUBLISHED_PRECISION = Decimal("0.001")


def _decimal(value: object) -> Decimal:
    try:
        return Decimal(str(value)).quantize(
            PUBLISHED_PRECISION,
            rounding=ROUND_HALF_UP,
        )
    except Exception as exc:
        raise ValueError(f"Invalid reference objective: {value!r}") from exc


def load_references(
    workbook_path: str | Path,
) -> dict[str, dict[str, dict[str, object]]]:
    """Load immutable objective-only Gil and Best Known benchmarks."""
    path = Path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"{path}: missing sheet {SHEET_NAME!r}")

    references: dict[str, dict[str, dict[str, object]]] = {}
    worksheet = workbook[SHEET_NAME]
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=5, max_col=5, values_only=True),
        start=5,
    ):
        if not values or not values[0]:
            continue
        instance_id = str(values[0]).strip()
        if not INSTANCE_RE.fullmatch(instance_id):
            raise ValueError(
                f"{path}:{row_number}: invalid instance id {instance_id!r}"
            )
        if instance_id in references:
            raise ValueError(f"{path}: duplicate instance id {instance_id}")
        if any(value is None for value in values[1:5]):
            raise ValueError(f"{path}:{row_number}: incomplete objective row")
        references[instance_id] = {
            "gil_2020_grasp_vnd": {
                "completion_time_s": _decimal(values[1]),
                "max_turnover_time_s": _decimal(values[2]),
                "source": f"{path.name}/{SHEET_NAME}",
                "benchmark_only": True,
            },
            "best_known": {
                "completion_time_s": _decimal(values[3]),
                "max_turnover_time_s": _decimal(values[4]),
                "source": f"{path.name}/{SHEET_NAME}",
                "benchmark_only": True,
            },
        }
    if len(references) != 64:
        raise ValueError(
            f"{path}: expected 64 reference rows, found {len(references)}"
        )
    return references


def normalize_actual(
    instance_id: str,
    algorithm_id: str,
    completed_tours: list[tuple],
    arrivals: dict[int, float],
) -> dict[str, object]:
    batches = [
        {
            "tour_id": int(row[0]),
            "route_start_time_s": float(row[1]),
            "completion_time_s": float(row[2]),
            "order_ids": sorted(int(value) for value in row[3]),
            "picker_id": int(row[4]),
        }
        for row in completed_tours
    ]
    completion = max(
        (batch["completion_time_s"] for batch in batches),
        default=0.0,
    )
    turnovers = [
        batch["completion_time_s"] - arrivals[order_id]
        for batch in batches
        for order_id in batch["order_ids"]
    ]
    return {
        "instance_id": instance_id,
        "algorithm_id": algorithm_id,
        "completion_time_s": completion,
        "max_turnover_time_s": max(turnovers, default=0.0),
        "batches": batches,
    }


def _compare_metric(
    actual: float,
    reference: Decimal,
    absolute_tolerance_s: float,
    relative_tolerance: float,
) -> dict[str, object]:
    reference_float = float(reference)
    gap = actual - reference_float
    equal = math.isclose(
        actual,
        reference_float,
        abs_tol=absolute_tolerance_s,
        rel_tol=relative_tolerance,
    )
    return {
        "actual": actual,
        "reference": str(reference),
        "signed_gap": gap,
        "absolute_gap": abs(gap),
        "relative_gap": (
            gap / reference_float if reference_float else math.inf
        ),
        "status": "EQUAL" if equal else ("BETTER" if gap < 0 else "WORSE"),
    }


def compare_benchmarks(
    actual: dict[str, object],
    references: dict[str, dict[str, object]],
    absolute_tolerance_s: float,
    relative_tolerance: float,
) -> list[dict[str, object]]:
    reports = []
    for algorithm_id, reference in references.items():
        completion = _compare_metric(
            float(actual["completion_time_s"]),
            reference["completion_time_s"],
            absolute_tolerance_s,
            relative_tolerance,
        )
        turnover = _compare_metric(
            float(actual["max_turnover_time_s"]),
            reference["max_turnover_time_s"],
            absolute_tolerance_s,
            relative_tolerance,
        )
        first = None
        if completion["status"] != "EQUAL":
            first = "completion_time_s"
        elif turnover["status"] != "EQUAL":
            first = "max_turnover_time_s"
        reports.append(
            {
                "instance_id": actual["instance_id"],
                "actual_algorithm_id": actual["algorithm_id"],
                "reference_algorithm_id": algorithm_id,
                "source": reference["source"],
                "benchmark_only": True,
                "completion_time": completion,
                "max_turnover_time": turnover,
                "batch_membership": "NOT_AVAILABLE",
                "dispatch_times": "NOT_AVAILABLE",
                "route_sequence": "NOT_AVAILABLE",
                "first_observable_difference": first,
            }
        )
    return reports
